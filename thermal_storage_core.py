#!/usr/bin/env python3
"""Core calculation utilities for solar heating storage sizing."""

from __future__ import annotations

from pathlib import Path
from typing import Callable

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

RHO_WATER = 1000.0  # kg/m3
CP_WATER = 4182.0  # J/(kg*K)

TANK_TYPE_DEFAULTS = {
    "short_term": {
        "label": "Short / Mid-Term Storage",
        "ratio_l_per_m2": 200.0,
        "loss_percent_per_day": 1.0,
        "cost_coeff": 1400.0,
        "cost_exponent": 0.78,
    },
    "seasonal": {
        "label": "Long-Term / Seasonal Storage",
        "ratio_l_per_m2": 2000.0,
        "loss_percent_per_day": 0.15,
        "cost_coeff": 900.0,
        "cost_exponent": 0.82,
    },
}

TANK_GEOMETRY_DEFAULTS = {
    "short_term": {
        "height_m": 2.4,
        "static_mixing_coeff": 0.020,
        "charge_mixing_coeff": 0.080,
        "discharge_mixing_coeff": 0.060,
    },
    "seasonal": {
        "height_m": 7.5,
        "static_mixing_coeff": 0.010,
        "charge_mixing_coeff": 0.045,
        "discharge_mixing_coeff": 0.035,
    },
}

INPUT_COLUMN_ALIASES = {
    "outside_temp": [
        "outside_temp",
        "outside_air_temp",
        "outsideairtemp",
        "outsideairtempc",
    ],
    "solar_rad": [
        "solar_rad",
        "solar_rad_55deg_surface",
        "solarrad55degsurface",
        "solarrad55degsurfacewhm2",
    ],
    "horizontal_solar_rad": [
        "horizontal_solar_rad",
        "horizontalsolarrad",
        "horizontalsolarradwhm2",
    ],
    "heating_load": ["heating_load", "heatingload", "heatingloadw", "heatingloadwh"],
    "timestamp": ["timestamp", "time_step", "timestep"],
}


def clamp(value: float, minimum: float, maximum: float) -> float:
    return max(minimum, min(value, maximum))


def parse_numeric_list(value: str) -> list[float]:
    items = []
    for part in value.split(","):
        text = part.strip()
        if text:
            items.append(float(text))
    if not items:
        raise ValueError("Please provide at least one numeric value.")
    return sorted(set(items))


def make_candidate_values(min_value: float, max_value: float, step: float) -> list[float]:
    count = int(round((max_value - min_value) / step))
    return [round(min_value + index * step, 10) for index in range(count + 1)]


def normalize_tank_types(tank_types: list[str] | tuple[str, ...] | None) -> list[str]:
    if not tank_types:
        return ["short_term"]
    normalized = []
    for tank_type in tank_types:
        if tank_type not in TANK_TYPE_DEFAULTS:
            raise ValueError(f"Unsupported tank type: {tank_type}")
        normalized.append(tank_type)
    return normalized


def tank_type_label(tank_type: str) -> str:
    return TANK_TYPE_DEFAULTS[tank_type]["label"]


def tank_type_ratio_m3_per_m2(
    tank_type: str,
    short_term_ratio_l_per_m2: float,
    seasonal_ratio_l_per_m2: float,
) -> float:
    if tank_type == "short_term":
        return short_term_ratio_l_per_m2 / 1000.0
    return seasonal_ratio_l_per_m2 / 1000.0


def tank_loss_percent_per_day(
    tank_type: str,
    short_term_loss_percent_per_day: float,
    seasonal_loss_percent_per_day: float,
) -> float:
    if tank_type == "short_term":
        return short_term_loss_percent_per_day
    return seasonal_loss_percent_per_day


def tank_cost_curve(
    tank_type: str,
    short_term_cost_coeff: float,
    short_term_cost_exponent: float,
    seasonal_cost_coeff: float,
    seasonal_cost_exponent: float,
) -> tuple[float, float]:
    if tank_type == "short_term":
        return short_term_cost_coeff, short_term_cost_exponent
    return seasonal_cost_coeff, seasonal_cost_exponent


def tank_geometry_value(
    tank_type: str,
    short_term_value: float,
    seasonal_value: float,
) -> float:
    return short_term_value if tank_type == "short_term" else seasonal_value


def energy_per_m3_kwh(tank_min_temp: float, tank_max_temp: float) -> float:
    delta_t = tank_max_temp - tank_min_temp
    if delta_t <= 0:
        raise ValueError("Tank maximum temperature must be higher than tank minimum temperature.")
    return RHO_WATER * CP_WATER * delta_t / 3.6e6


def first_matching_column(df: pd.DataFrame, candidates: list[str]) -> str | None:
    normalized = {normalize_column_name(column): column for column in df.columns}
    for candidate in candidates:
        if candidate in df.columns:
            return candidate
        normalized_candidate = normalize_column_name(candidate)
        if normalized_candidate in normalized:
            return normalized[normalized_candidate]
    return None


def normalize_column_name(name: str) -> str:
    normalized = str(name).strip().lower()
    replacements = {
        "℃": "c",
        "°c": "c",
        "°": "",
        "（": "(",
        "）": ")",
    }
    for source, target in replacements.items():
        normalized = normalized.replace(source, target)
    return "".join(character for character in normalized if character.isalnum())


def load_input(
    filename: str | Path,
    radiation_unit: str = "kwh/m2",
    load_unit: str = "kwh",
    timestep_hours: float = 1.0,
) -> pd.DataFrame:
    file_path = Path(filename)
    suffix = file_path.suffix.lower()
    if suffix in {".xlsx", ".xls"}:
        df = pd.read_excel(file_path)
    else:
        df = pd.read_csv(file_path)
    mapped_columns: dict[str, str] = {}
    for standard_name, aliases in INPUT_COLUMN_ALIASES.items():
        found = first_matching_column(df, aliases)
        if found is not None:
            mapped_columns[standard_name] = found

    required = ["outside_temp", "solar_rad", "heating_load"]
    missing = [column for column in required if column not in mapped_columns]
    if missing:
        raise ValueError(f"Missing required columns in {filename}: {missing}")

    df = df.copy()
    df["outside_temp"] = df[mapped_columns["outside_temp"]].astype(float)
    df["solar_rad"] = df[mapped_columns["solar_rad"]].astype(float)
    df["heating_load"] = df[mapped_columns["heating_load"]].astype(float)
    if "horizontal_solar_rad" in mapped_columns:
        df["horizontal_solar_rad"] = df[mapped_columns["horizontal_solar_rad"]].astype(float)
        horizontal_column_name = normalize_column_name(mapped_columns["horizontal_solar_rad"])
        if radiation_unit.lower() == "w/m2":
            df["horizontal_solar_rad"] = df["horizontal_solar_rad"] * timestep_hours / 1000.0
        elif "whm2" in horizontal_column_name:
            df["horizontal_solar_rad"] = df["horizontal_solar_rad"] / 1000.0

    solar_column_name = normalize_column_name(mapped_columns["solar_rad"])
    if radiation_unit.lower() == "w/m2":
        df["solar_rad"] = df["solar_rad"] * timestep_hours / 1000.0
    elif "whm2" in solar_column_name or (suffix in {".xlsx", ".xls"} and df["solar_rad"].max() > 50.0):
        df["solar_rad"] = df["solar_rad"] / 1000.0
    elif radiation_unit.lower() != "kwh/m2":
        raise ValueError("Unsupported radiation unit, choose kWh/m2 or W/m2.")

    heating_column_name = normalize_column_name(mapped_columns["heating_load"])
    if load_unit.lower() == "w":
        df["heating_load"] = df["heating_load"] * timestep_hours / 1000.0
    elif heating_column_name.endswith("w") or heating_column_name.endswith("wh"):
        df["heating_load"] = df["heating_load"] * timestep_hours / 1000.0
    elif load_unit.lower() != "kwh":
        raise ValueError("Unsupported load unit, choose kWh or W.")

    if "timestamp" in mapped_columns:
        df["timestamp"] = pd.to_datetime(df[mapped_columns["timestamp"]])
    else:
        df["timestamp"] = pd.date_range(
            start="2025-01-01 00:00:00",
            periods=len(df),
            freq=pd.to_timedelta(timestep_hours, unit="h"),
        )

    return df


def collector_output_kwh(
    area_m2: float,
    solar_rad_kwh_m2: float,
    outside_temp_c: float,
    fluid_temp_c: float,
    timestep_hours: float,
    eta0: float,
    a1: float,
    a2: float,
    collector_soiling_factor: float = 1.0,
) -> tuple[float, float]:
    if solar_rad_kwh_m2 <= 0 or area_m2 <= 0:
        return 0.0, 0.0

    irradiance_w_m2 = solar_rad_kwh_m2 * 1000.0 / max(timestep_hours, 1e-9)
    delta_t = fluid_temp_c - outside_temp_c
    eta = eta0 - a1 * delta_t / max(irradiance_w_m2, 1e-6) - a2 * delta_t**2 / max(irradiance_w_m2, 1e-6)
    eta = clamp(eta, 0.0, 0.85) * clamp(collector_soiling_factor, 0.0, 1.0)
    collected_kwh = area_m2 * solar_rad_kwh_m2 * eta
    return collected_kwh, eta


def dynamic_collector_delta_t(
    area_m2: float,
    solar_rad_kwh_m2: float,
    outside_temp_c: float,
    inlet_temp_c: float,
    timestep_hours: float,
    eta0: float,
    a1: float,
    a2: float,
    collector_soiling_factor: float,
    nominal_delta_t_c: float,
    collector_flow_lph_per_m2: float,
    collector_delta_t_min_c: float,
    collector_delta_t_max_c: float,
    max_outlet_temp_c: float,
    iterations: int = 3,
) -> tuple[float, float, float]:
    if solar_rad_kwh_m2 <= 0 or area_m2 <= 0:
        return 0.0, 0.0, 0.0

    mass_flow_kg_s = max(area_m2 * collector_flow_lph_per_m2 / 3600.0, 1e-9)
    delta_t_guess_c = clamp(nominal_delta_t_c, collector_delta_t_min_c, collector_delta_t_max_c)
    collected_kwh = 0.0
    efficiency = 0.0

    for _ in range(max(iterations, 1)):
        outlet_temp_c = min(inlet_temp_c + delta_t_guess_c, max_outlet_temp_c)
        mean_temp_c = (inlet_temp_c + outlet_temp_c) / 2.0
        collected_kwh, efficiency = collector_output_kwh(
            area_m2=area_m2,
            solar_rad_kwh_m2=solar_rad_kwh_m2,
            outside_temp_c=outside_temp_c,
            fluid_temp_c=mean_temp_c,
            timestep_hours=timestep_hours,
            eta0=eta0,
            a1=a1,
            a2=a2,
            collector_soiling_factor=collector_soiling_factor,
        )
        useful_power_w = collected_kwh * 1000.0 / max(timestep_hours, 1e-9)
        calculated_delta_t_c = useful_power_w / max(mass_flow_kg_s * CP_WATER, 1e-9)
        max_physical_delta_t_c = max(max_outlet_temp_c - inlet_temp_c, 0.0)
        delta_t_guess_c = clamp(
            calculated_delta_t_c,
            min(collector_delta_t_min_c, max_physical_delta_t_c),
            min(collector_delta_t_max_c, max_physical_delta_t_c),
        )
        if max_physical_delta_t_c <= 0:
            delta_t_guess_c = 0.0
            break

    return delta_t_guess_c, collected_kwh, efficiency


def temperature_fraction(temp_c: float, tank_min_temp: float, tank_max_temp: float) -> float:
    return clamp((temp_c - tank_min_temp) / max(tank_max_temp - tank_min_temp, 1e-9), 0.0, 1.0)


def storage_energy_from_temp(
    tank_temp_c: float,
    storage_capacity_kwh: float,
    tank_min_temp: float,
    tank_max_temp: float,
) -> float:
    return storage_capacity_kwh * temperature_fraction(tank_temp_c, tank_min_temp, tank_max_temp)


def storage_temp_from_energy(
    storage_energy_kwh: float,
    storage_capacity_kwh: float,
    tank_min_temp: float,
    tank_max_temp: float,
) -> float:
    if storage_capacity_kwh <= 0:
        return tank_min_temp
    fraction = clamp(storage_energy_kwh / storage_capacity_kwh, 0.0, 1.0)
    return tank_min_temp + fraction * (tank_max_temp - tank_min_temp)


def energy_above_temperature(
    storage_energy_kwh: float,
    storage_capacity_kwh: float,
    tank_min_temp: float,
    tank_max_temp: float,
    threshold_temp_c: float,
) -> float:
    threshold_energy = storage_energy_from_temp(
        tank_temp_c=threshold_temp_c,
        storage_capacity_kwh=storage_capacity_kwh,
        tank_min_temp=tank_min_temp,
        tank_max_temp=tank_max_temp,
    )
    return max(storage_energy_kwh - threshold_energy, 0.0)


def energy_between_temperatures(
    storage_energy_kwh: float,
    storage_capacity_kwh: float,
    tank_min_temp: float,
    tank_max_temp: float,
    low_temp_c: float,
    high_temp_c: float,
) -> float:
    if high_temp_c <= low_temp_c:
        return 0.0
    energy_above_low = energy_above_temperature(
        storage_energy_kwh=storage_energy_kwh,
        storage_capacity_kwh=storage_capacity_kwh,
        tank_min_temp=tank_min_temp,
        tank_max_temp=tank_max_temp,
        threshold_temp_c=low_temp_c,
    )
    energy_above_high = energy_above_temperature(
        storage_energy_kwh=storage_energy_kwh,
        storage_capacity_kwh=storage_capacity_kwh,
        tank_min_temp=tank_min_temp,
        tank_max_temp=tank_max_temp,
        threshold_temp_c=high_temp_c,
    )
    return max(energy_above_low - energy_above_high, 0.0)


def storage_loss_kwh(storage_energy_kwh: float, loss_percent_per_day: float, timestep_hours: float) -> float:
    daily_retention = float(np.clip(1.0 - loss_percent_per_day / 100.0, 0.0, 1.0))
    step_retention = daily_retention ** (timestep_hours / 24.0)
    return storage_energy_kwh * (1.0 - step_retention)


def default_stratification_nodes(tank_type: str) -> int:
    return 8 if tank_type == "short_term" else 12


def tank_diameter_from_volume_height(tank_volume_m3: float, tank_height_m: float) -> float:
    if tank_volume_m3 <= 0 or tank_height_m <= 0:
        raise ValueError("Tank volume and tank height must both be greater than zero.")
    return float(np.sqrt(4.0 * tank_volume_m3 / (np.pi * tank_height_m)))


def stratification_efficiency_from_geometry(tank_height_m: float, tank_diameter_m: float) -> float:
    aspect_ratio = tank_height_m / max(tank_diameter_m, 1e-9)
    aspect_factor = clamp((aspect_ratio - 0.9) / 2.6, 0.0, 1.0)
    height_factor = clamp((tank_height_m - 1.5) / 5.0, 0.0, 1.0)
    return clamp(0.35 + 0.45 * aspect_factor + 0.15 * height_factor, 0.25, 0.95)


def adjusted_mixing_coefficient(base_coeff: float, stratification_efficiency: float) -> float:
    geometry_penalty = 1.25 - stratification_efficiency
    return clamp(base_coeff * geometry_penalty, 0.0, 0.45)


def mix_tank_layers(layer_temps_c: np.ndarray, mixing_coeff: float) -> np.ndarray:
    coeff = clamp(mixing_coeff, 0.0, 0.45)
    if coeff <= 0 or len(layer_temps_c) <= 1:
        return stabilize_stratification(layer_temps_c.copy())

    previous = layer_temps_c.copy()
    updated = previous.copy()
    updated[0] = previous[0] + coeff * (previous[1] - previous[0])
    updated[-1] = previous[-1] + coeff * (previous[-2] - previous[-1])
    if len(previous) > 2:
        updated[1:-1] = previous[1:-1] + coeff * (
            previous[:-2] - 2.0 * previous[1:-1] + previous[2:]
        )
    return stabilize_stratification(updated)


def tank_geometry_metrics(
    tank_volume_m3: float,
    tank_height_m: float,
    static_mixing_coeff: float,
    charge_mixing_coeff: float,
    discharge_mixing_coeff: float,
) -> dict[str, float]:
    tank_diameter_m = tank_diameter_from_volume_height(tank_volume_m3, tank_height_m)
    height_to_diameter_ratio = tank_height_m / max(tank_diameter_m, 1e-9)
    ideal_stratification = bool(
        np.isclose(static_mixing_coeff, 0.0)
        and np.isclose(charge_mixing_coeff, 0.0)
        and np.isclose(discharge_mixing_coeff, 0.0)
    )
    stratification_efficiency = (
        1.0
        if ideal_stratification
        else stratification_efficiency_from_geometry(
            tank_height_m=tank_height_m,
            tank_diameter_m=tank_diameter_m,
        )
    )
    return {
        "tank_height_m": float(tank_height_m),
        "tank_diameter_m": tank_diameter_m,
        "tank_height_to_diameter_ratio": float(height_to_diameter_ratio),
        "stratification_efficiency": stratification_efficiency,
        "ideal_stratification": float(ideal_stratification),
        "static_mixing_coeff_effective": adjusted_mixing_coefficient(
            static_mixing_coeff,
            stratification_efficiency,
        ),
        "charge_mixing_coeff_effective": adjusted_mixing_coefficient(
            charge_mixing_coeff,
            stratification_efficiency,
        ),
        "discharge_mixing_coeff_effective": adjusted_mixing_coefficient(
            discharge_mixing_coeff,
            stratification_efficiency,
        ),
    }


def create_tank_layers(initial_temp: float, num_nodes: int) -> np.ndarray:
    return np.full(num_nodes, float(initial_temp), dtype=float)


def stabilize_stratification(layer_temps_c: np.ndarray) -> np.ndarray:
    return np.sort(layer_temps_c)[::-1]


def layer_kwh_per_k(tank_volume_m3: float, num_nodes: int) -> float:
    return tank_volume_m3 * RHO_WATER * CP_WATER / 3.6e6 / max(num_nodes, 1)


def tank_energy_from_layers(layer_temps_c: np.ndarray, tank_min_temp: float, kwh_per_k: float) -> float:
    return float(np.maximum(layer_temps_c - tank_min_temp, 0.0).sum() * kwh_per_k)


def tank_energy_quality_breakdown(
    layer_temps_c: np.ndarray,
    tank_min_temp: float,
    heating_return_temp: float,
    heating_supply_temp: float,
    kwh_per_k: float,
) -> tuple[float, float, float]:
    total_stored_kwh = tank_energy_from_layers(layer_temps_c, tank_min_temp, kwh_per_k)
    above_return_kwh = float(np.maximum(layer_temps_c - heating_return_temp, 0.0).sum() * kwh_per_k)
    above_supply_kwh = float(np.maximum(layer_temps_c - heating_supply_temp, 0.0).sum() * kwh_per_k)
    between_return_and_supply_kwh = max(above_return_kwh - above_supply_kwh, 0.0)
    below_return_kwh = max(total_stored_kwh - above_return_kwh, 0.0)
    return above_supply_kwh, between_return_and_supply_kwh, below_return_kwh


def tank_summary_temperatures(layer_temps_c: np.ndarray) -> tuple[float, float, float]:
    return float(layer_temps_c.mean()), float(layer_temps_c[0]), float(layer_temps_c[-1])


def charge_tank_layers(
    layer_temps_c: np.ndarray,
    charge_kwh: float,
    source_temp_c: float,
    tank_max_temp: float,
    kwh_per_k: float,
) -> tuple[np.ndarray, float]:
    remaining_kwh = max(charge_kwh, 0.0)
    updated = layer_temps_c.copy()

    for idx in range(len(updated)):
        if remaining_kwh <= 0:
            break
        target_temp = min(tank_max_temp, source_temp_c)
        if target_temp <= updated[idx]:
            continue
        capacity_kwh = (target_temp - updated[idx]) * kwh_per_k
        applied_kwh = min(remaining_kwh, capacity_kwh)
        updated[idx] += applied_kwh / kwh_per_k
        remaining_kwh -= applied_kwh

    return stabilize_stratification(updated), charge_kwh - remaining_kwh


def discharge_tank_above_threshold(
    layer_temps_c: np.ndarray,
    demand_kwh: float,
    threshold_temp_c: float,
    kwh_per_k: float,
) -> tuple[np.ndarray, float]:
    remaining_kwh = max(demand_kwh, 0.0)
    updated = layer_temps_c.copy()
    delivered_kwh = 0.0

    for idx in range(len(updated)):
        if remaining_kwh <= 0:
            break
        available_kwh = max(updated[idx] - threshold_temp_c, 0.0) * kwh_per_k
        used_kwh = min(remaining_kwh, available_kwh)
        updated[idx] -= used_kwh / kwh_per_k
        remaining_kwh -= used_kwh
        delivered_kwh += used_kwh

    return stabilize_stratification(updated), delivered_kwh


def discharge_tank_to_return_with_reinjection(
    layer_temps_c: np.ndarray,
    demand_kwh: float,
    supply_temp_c: float,
    return_temp_c: float,
    tank_min_temp: float,
    kwh_per_k: float,
) -> tuple[np.ndarray, float]:
    remaining_kwh = max(demand_kwh, 0.0)
    updated = layer_temps_c.copy()
    delivered_kwh = 0.0
    effective_return_temp_c = max(return_temp_c, tank_min_temp)

    for idx in range(len(updated)):
        if remaining_kwh <= 0:
            break
        node_temp_c = updated[idx]
        if node_temp_c <= supply_temp_c:
            continue

        available_kwh = max(node_temp_c - effective_return_temp_c, 0.0) * kwh_per_k
        used_kwh = min(remaining_kwh, available_kwh)
        if used_kwh <= 0:
            continue

        circulation_fraction = used_kwh / max(available_kwh, 1e-9)
        bottom_temp_before_c = updated[-1]
        bottom_temp_after_c = bottom_temp_before_c + circulation_fraction * (
            effective_return_temp_c - bottom_temp_before_c
        )
        bottom_energy_change_kwh = (bottom_temp_after_c - bottom_temp_before_c) * kwh_per_k

        updated[-1] = bottom_temp_after_c
        updated[idx] -= (used_kwh + bottom_energy_change_kwh) / kwh_per_k
        updated[idx] = max(updated[idx], tank_min_temp)

        remaining_kwh -= used_kwh
        delivered_kwh += used_kwh

    return stabilize_stratification(updated), delivered_kwh


def discharge_tank_for_hp_source(
    layer_temps_c: np.ndarray,
    source_energy_needed_kwh: float,
    min_source_temp_c: float,
    max_source_temp_c: float,
    kwh_per_k: float,
) -> tuple[np.ndarray, float]:
    remaining_kwh = max(source_energy_needed_kwh, 0.0)
    updated = layer_temps_c.copy()
    extracted_kwh = 0.0

    for idx in range(len(updated)):
        if remaining_kwh <= 0:
            break
        upper_temp = min(updated[idx], max_source_temp_c)
        lower_temp = min_source_temp_c
        if upper_temp <= lower_temp:
            continue
        available_kwh = (upper_temp - lower_temp) * kwh_per_k
        used_kwh = min(remaining_kwh, available_kwh)
        updated[idx] -= used_kwh / kwh_per_k
        remaining_kwh -= used_kwh
        extracted_kwh += used_kwh

    return stabilize_stratification(updated), extracted_kwh


def apply_tank_losses(
    layer_temps_c: np.ndarray,
    tank_min_temp: float,
    loss_percent_per_day: float,
    timestep_hours: float,
    kwh_per_k: float,
) -> tuple[np.ndarray, float]:
    daily_retention = clamp(1.0 - loss_percent_per_day / 100.0, 0.0, 1.0)
    step_retention = daily_retention ** (timestep_hours / 24.0)
    updated = tank_min_temp + (layer_temps_c - tank_min_temp) * step_retention
    loss_kwh = float(np.maximum(layer_temps_c - updated, 0.0).sum() * kwh_per_k)
    return stabilize_stratification(updated), loss_kwh


def tank_capital_cost(
    tank_volume_m3: float,
    tank_type: str,
    short_term_cost_coeff: float,
    short_term_cost_exponent: float,
    seasonal_cost_coeff: float,
    seasonal_cost_exponent: float,
) -> float:
    coeff, exponent = tank_cost_curve(
        tank_type=tank_type,
        short_term_cost_coeff=short_term_cost_coeff,
        short_term_cost_exponent=short_term_cost_exponent,
        seasonal_cost_coeff=seasonal_cost_coeff,
        seasonal_cost_exponent=seasonal_cost_exponent,
    )
    return coeff * max(tank_volume_m3, 0.0) ** exponent


def _simulate_single_pass(
    weather_load_df: pd.DataFrame,
    collector_area_m2: float,
    tank_layers_c: np.ndarray,
    tank_type: str,
    tank_min_temp: float,
    tank_max_temp: float,
    heating_supply_temp: float,
    heating_delta_t: float,
    solar_collector_delta_t: float,
    collector_flow_lph_per_m2: float,
    collector_delta_t_min_c: float,
    collector_delta_t_max_c: float,
    direct_solar_delivery_margin_c: float,
    tank_thermal_loss_percent_per_day: float,
    auxiliary_system: str,
    auxiliary_efficiency: float,
    backup_auxiliary_efficiency: float,
    timestep_hours: float,
    eta0: float,
    a1: float,
    a2: float,
    collector_soiling_factor: float,
    kwh_per_k_per_node: float,
    tank_geometry: dict[str, float],
    pass_index: int,
    return_hourly: bool,
) -> tuple[dict[str, float], pd.DataFrame, np.ndarray]:
    heating_return_temp = heating_supply_temp - heating_delta_t
    working_layers = tank_layers_c.copy()
    records: list[dict[str, float]] | None = [] if return_hourly else None
    totals = {
        "solar_energy_collected_kwh": 0.0,
        "solar_direct_collector_to_load_kwh": 0.0,
        "solar_from_tank_direct_kwh": 0.0,
        "solar_from_tank_via_hp_kwh": 0.0,
        "solar_from_tank_via_hp_source_kwh": 0.0,
        "solar_source_energy_used_kwh": 0.0,
        "solar_energy_for_heating_kwh": 0.0,
        "auxiliary_heating_delivered_kwh": 0.0,
        "auxiliary_energy_input_kwh": 0.0,
        "auxiliary_energy_for_hp_boost_kwh": 0.0,
        "solar_energy_curtailed_kwh": 0.0,
        "tank_charge_kwh": 0.0,
        "tank_discharge_total_kwh": 0.0,
        "tank_loss_kwh": 0.0,
        "tank_energy_above_supply_kwh_hours": 0.0,
        "tank_energy_between_return_supply_kwh_hours": 0.0,
        "tank_energy_below_return_kwh_hours": 0.0,
        "heating_load_kwh": float(weather_load_df["heating_load"].sum()),
        "_productive_efficiency_sum": 0.0,
        "_productive_efficiency_count": 0.0,
        "_final_tank_temp_c": float(tank_layers_c.mean()) if len(tank_layers_c) else tank_min_temp,
    }

    for row in weather_load_df.itertuples(index=False):
        tank_temp_start, tank_top_temp_start, tank_bottom_temp_start = tank_summary_temperatures(working_layers)
        storage_energy_start_kwh = tank_energy_from_layers(working_layers, tank_min_temp, kwh_per_k_per_node)
        collector_enabled = not (tank_type == "short_term" and row.heating_load <= 0)
        if not collector_enabled:
            solar_collected_kwh = 0.0
            efficiency = 0.0
            direct_solar_to_load_kwh = 0.0
            remaining_load_kwh = row.heating_load
            surplus_solar_kwh = 0.0
            tank_charge_kwh = 0.0
            solar_curtailed_kwh = 0.0
            direct_mode_enabled = False
            charge_source_temp = tank_temp_start
            collector_inlet_temp = max(tank_temp_start, tank_min_temp)
            collector_outlet_temp = collector_inlet_temp
            actual_collector_delta_t_c = 0.0
        else:
            if tank_type == "short_term":
                collector_inlet_reference_temp = tank_temp_start
            elif row.heating_load > 0:
                collector_inlet_reference_temp = max(tank_bottom_temp_start, heating_return_temp)
            else:
                collector_inlet_reference_temp = tank_bottom_temp_start
            collector_inlet_temp = max(collector_inlet_reference_temp, tank_min_temp)
            actual_collector_delta_t_c, solar_collected_kwh, efficiency = dynamic_collector_delta_t(
                area_m2=collector_area_m2,
                solar_rad_kwh_m2=row.solar_rad,
                outside_temp_c=row.outside_temp,
                inlet_temp_c=collector_inlet_temp,
                timestep_hours=timestep_hours,
                eta0=eta0,
                a1=a1,
                a2=a2,
                collector_soiling_factor=collector_soiling_factor,
                nominal_delta_t_c=solar_collector_delta_t,
                collector_flow_lph_per_m2=collector_flow_lph_per_m2,
                collector_delta_t_min_c=collector_delta_t_min_c,
                collector_delta_t_max_c=collector_delta_t_max_c,
                max_outlet_temp_c=tank_max_temp,
            )
            collector_outlet_temp = min(collector_inlet_temp + actual_collector_delta_t_c, tank_max_temp)
            # Short-term storage is modeled as tank-first: solar charges storage before
            # any heat is delivered to the load from the tank in the same timestep.
            if tank_type == "short_term":
                direct_mode_enabled = False
                direct_solar_to_load_kwh = 0.0
            else:
                direct_mode_enabled = (
                    row.heating_load > 0
                    and solar_collected_kwh > 0
                    and collector_outlet_temp >= heating_supply_temp
                )
                if direct_mode_enabled:
                    direct_solar_to_load_kwh = min(solar_collected_kwh, row.heating_load)
                else:
                    direct_solar_to_load_kwh = 0.0
            charge_source_temp = collector_outlet_temp
            remaining_load_kwh = row.heating_load - direct_solar_to_load_kwh

            surplus_solar_kwh = max(solar_collected_kwh - direct_solar_to_load_kwh, 0.0)
            available_storage_kwh = max(
                tank_volume_kwh_capacity_from_layers(working_layers, tank_min_temp, tank_max_temp, kwh_per_k_per_node)
                - storage_energy_start_kwh,
                0.0,
            )
            candidate_charge_kwh = min(surplus_solar_kwh, available_storage_kwh)
            working_layers, tank_charge_kwh = charge_tank_layers(
                layer_temps_c=working_layers,
                charge_kwh=candidate_charge_kwh,
                source_temp_c=charge_source_temp,
                tank_max_temp=tank_max_temp,
                kwh_per_k=kwh_per_k_per_node,
            )
            if tank_charge_kwh > 0:
                charge_fraction = tank_charge_kwh / max(candidate_charge_kwh, 1e-9)
                working_layers = mix_tank_layers(
                    working_layers,
                    tank_geometry["charge_mixing_coeff_effective"] * charge_fraction,
                )
            solar_curtailed_kwh = surplus_solar_kwh - tank_charge_kwh
        working_layers, solar_from_tank_direct_kwh = discharge_tank_to_return_with_reinjection(
            layer_temps_c=working_layers,
            demand_kwh=remaining_load_kwh,
            supply_temp_c=heating_supply_temp,
            return_temp_c=heating_return_temp,
            tank_min_temp=tank_min_temp,
            kwh_per_k=kwh_per_k_per_node,
        )
        if solar_from_tank_direct_kwh > 0:
            draw_fraction = solar_from_tank_direct_kwh / max(row.heating_load, 1e-9)
            working_layers = mix_tank_layers(
                working_layers,
                tank_geometry["discharge_mixing_coeff_effective"] * draw_fraction,
            )
        remaining_load_kwh -= solar_from_tank_direct_kwh

        solar_from_tank_via_hp_kwh = 0.0
        solar_from_tank_via_hp_source_kwh = 0.0
        auxiliary_energy_for_hp_boost_kwh = 0.0
        if remaining_load_kwh > 0 and auxiliary_system == "water_heat_pump" and auxiliary_efficiency > 1.0:
            hp_deliverable_factor = auxiliary_efficiency / max(auxiliary_efficiency - 1.0, 1e-6)
            hp_source_needed_kwh = remaining_load_kwh / hp_deliverable_factor
            working_layers, solar_from_tank_via_hp_source_kwh = discharge_tank_for_hp_source(
                layer_temps_c=working_layers,
                source_energy_needed_kwh=hp_source_needed_kwh,
                min_source_temp_c=tank_min_temp,
                max_source_temp_c=min(heating_supply_temp, tank_max_temp),
                kwh_per_k=kwh_per_k_per_node,
            )
            if solar_from_tank_via_hp_source_kwh > 0:
                solar_from_tank_via_hp_kwh = solar_from_tank_via_hp_source_kwh * hp_deliverable_factor
                solar_from_tank_via_hp_kwh = min(remaining_load_kwh, solar_from_tank_via_hp_kwh)
                solar_from_tank_via_hp_source_kwh = solar_from_tank_via_hp_kwh * (1.0 - 1.0 / auxiliary_efficiency)
                auxiliary_energy_for_hp_boost_kwh = solar_from_tank_via_hp_kwh / auxiliary_efficiency
                hp_draw_fraction = solar_from_tank_via_hp_source_kwh / max(hp_source_needed_kwh, 1e-9)
                working_layers = mix_tank_layers(
                    working_layers,
                    tank_geometry["discharge_mixing_coeff_effective"] * hp_draw_fraction,
                )
                remaining_load_kwh -= solar_from_tank_via_hp_kwh

        auxiliary_heating_delivered_kwh = max(remaining_load_kwh, 0.0)
        if auxiliary_system == "water_heat_pump" and auxiliary_efficiency > 1.0:
            auxiliary_energy_input_kwh = (
                auxiliary_heating_delivered_kwh / max(backup_auxiliary_efficiency, 1e-6)
                + auxiliary_energy_for_hp_boost_kwh
            )
        else:
            auxiliary_energy_input_kwh = (
                auxiliary_heating_delivered_kwh / max(auxiliary_efficiency, 1e-6) + auxiliary_energy_for_hp_boost_kwh
            )

        working_layers, tank_loss_step_kwh = apply_tank_losses(
            layer_temps_c=working_layers,
            tank_min_temp=tank_min_temp,
            loss_percent_per_day=tank_thermal_loss_percent_per_day,
            timestep_hours=timestep_hours,
            kwh_per_k=kwh_per_k_per_node,
        )
        working_layers = mix_tank_layers(working_layers, tank_geometry["static_mixing_coeff_effective"])
        tank_temp_end, tank_top_temp_end, tank_bottom_temp_end = tank_summary_temperatures(working_layers)
        storage_energy_kwh = tank_energy_from_layers(working_layers, tank_min_temp, kwh_per_k_per_node)
        tank_energy_above_supply_kwh, tank_energy_between_return_supply_kwh, tank_energy_below_return_kwh = (
            tank_energy_quality_breakdown(
                layer_temps_c=working_layers,
                tank_min_temp=tank_min_temp,
                heating_return_temp=heating_return_temp,
                heating_supply_temp=heating_supply_temp,
                kwh_per_k=kwh_per_k_per_node,
            )
        )

        solar_heating_delivered_kwh = direct_solar_to_load_kwh + solar_from_tank_direct_kwh + solar_from_tank_via_hp_kwh
        solar_source_energy_used_kwh = (
            direct_solar_to_load_kwh + solar_from_tank_direct_kwh + solar_from_tank_via_hp_source_kwh
        )
        tank_discharge_total_kwh = solar_from_tank_direct_kwh + solar_from_tank_via_hp_source_kwh
        if solar_collected_kwh > 0:
            totals["_productive_efficiency_sum"] += efficiency
            totals["_productive_efficiency_count"] += 1.0
        totals["_final_tank_temp_c"] = tank_temp_end

        for key, value in {
            "solar_energy_collected_kwh": solar_collected_kwh,
            "solar_direct_collector_to_load_kwh": direct_solar_to_load_kwh,
            "solar_from_tank_direct_kwh": solar_from_tank_direct_kwh,
            "solar_from_tank_via_hp_kwh": solar_from_tank_via_hp_kwh,
            "solar_from_tank_via_hp_source_kwh": solar_from_tank_via_hp_source_kwh,
            "solar_source_energy_used_kwh": solar_source_energy_used_kwh,
            "solar_energy_for_heating_kwh": solar_heating_delivered_kwh,
            "auxiliary_heating_delivered_kwh": auxiliary_heating_delivered_kwh,
            "auxiliary_energy_input_kwh": auxiliary_energy_input_kwh,
            "auxiliary_energy_for_hp_boost_kwh": auxiliary_energy_for_hp_boost_kwh,
            "solar_energy_curtailed_kwh": solar_curtailed_kwh,
            "tank_charge_kwh": tank_charge_kwh,
            "tank_discharge_total_kwh": tank_discharge_total_kwh,
            "tank_loss_kwh": tank_loss_step_kwh,
            "tank_energy_above_supply_kwh_hours": tank_energy_above_supply_kwh * timestep_hours,
            "tank_energy_between_return_supply_kwh_hours": tank_energy_between_return_supply_kwh * timestep_hours,
            "tank_energy_below_return_kwh_hours": tank_energy_below_return_kwh * timestep_hours,
        }.items():
            totals[key] += value

        if records is not None:
            timestamp = row.timestamp
            if pass_index > 0:
                timestamp = pd.Timestamp(timestamp) + pd.DateOffset(years=pass_index)

            records.append(
                {
                    "timestamp": timestamp,
                    "simulation_year_index": pass_index + 1,
                    "tank_type": tank_type,
                    "tank_type_label": tank_type_label(tank_type),
                    "outside_temp_c": row.outside_temp,
                    "solar_rad_kwh_m2": row.solar_rad,
                    "horizontal_solar_rad_kwh_m2": getattr(row, "horizontal_solar_rad", np.nan),
                    "heating_load_kwh": row.heating_load,
                    "collector_efficiency": efficiency,
                    "solar_energy_collected_kwh": solar_collected_kwh,
                    "solar_direct_collector_to_load_kwh": direct_solar_to_load_kwh,
                    "solar_from_tank_direct_kwh": solar_from_tank_direct_kwh,
                    "solar_from_tank_via_hp_kwh": solar_from_tank_via_hp_kwh,
                    "solar_from_tank_via_hp_source_kwh": solar_from_tank_via_hp_source_kwh,
                    "solar_source_energy_used_kwh": solar_source_energy_used_kwh,
                    "solar_heating_total_kwh": solar_heating_delivered_kwh,
                    "auxiliary_heating_delivered_kwh": auxiliary_heating_delivered_kwh,
                    "auxiliary_energy_input_kwh": auxiliary_energy_input_kwh,
                    "auxiliary_energy_for_hp_boost_kwh": auxiliary_energy_for_hp_boost_kwh,
                    "tank_temp_start_c": tank_temp_start,
                    "tank_temp_end_c": tank_temp_end,
                    "tank_top_temp_start_c": tank_top_temp_start,
                    "tank_bottom_temp_start_c": tank_bottom_temp_start,
                    "tank_top_temp_end_c": tank_top_temp_end,
                    "tank_bottom_temp_end_c": tank_bottom_temp_end,
                    "tank_stored_energy_kwh": storage_energy_kwh,
                    "tank_charge_kwh": tank_charge_kwh,
                    "tank_discharge_direct_kwh": solar_from_tank_direct_kwh,
                    "tank_discharge_hp_source_kwh": solar_from_tank_via_hp_source_kwh,
                    "tank_discharge_total_kwh": tank_discharge_total_kwh,
                    "tank_loss_kwh": tank_loss_step_kwh,
                    "solar_energy_curtailed_kwh": solar_curtailed_kwh,
                    "tank_energy_above_supply_kwh": tank_energy_above_supply_kwh,
                    "tank_energy_between_return_supply_kwh": tank_energy_between_return_supply_kwh,
                    "tank_energy_below_return_kwh": tank_energy_below_return_kwh,
                    "tank_height_m": tank_geometry["tank_height_m"],
                    "tank_diameter_m": tank_geometry["tank_diameter_m"],
                    "tank_height_to_diameter_ratio": tank_geometry["tank_height_to_diameter_ratio"],
                    "stratification_efficiency": tank_geometry["stratification_efficiency"],
                    "collector_direct_mode_enabled": float(direct_mode_enabled),
                    "collector_inlet_temp_c": collector_inlet_temp,
                    "collector_outlet_temp_c": collector_outlet_temp,
                    "collector_delta_t_c": actual_collector_delta_t_c if collector_enabled else 0.0,
                }
            )

    records_df = pd.DataFrame.from_records(records) if records is not None else pd.DataFrame()
    return totals, records_df, working_layers


def tank_volume_kwh_capacity_from_layers(
    layer_temps_c: np.ndarray,
    tank_min_temp: float,
    tank_max_temp: float,
    kwh_per_k: float,
) -> float:
    return len(layer_temps_c) * kwh_per_k * (tank_max_temp - tank_min_temp)


def simulate_design(
    weather_load_df: pd.DataFrame,
    collector_area_m2: float,
    tank_volume_m3: float,
    tank_type: str = "short_term",
    tank_min_temp: float = 30.0,
    tank_max_temp: float = 80.0,
    initial_tank_temp: float | None = None,
    heating_supply_temp: float = 45.0,
    heating_delta_t: float = 10.0,
    solar_collector_delta_t: float = 10.0,
    collector_flow_lph_per_m2: float = 20.0,
    collector_delta_t_min_c: float = 2.0,
    collector_delta_t_max_c: float = 15.0,
    direct_solar_delivery_margin_c: float = 0.0,
    tank_thermal_loss_percent_per_day: float = 1.0,
    auxiliary_system: str = "electric",
    auxiliary_efficiency: float = 0.98,
    backup_auxiliary_efficiency: float = 0.98,
    timestep_hours: float = 1.0,
    eta0: float = 0.72,
    a1: float = 3.2,
    a2: float = 0.0,
    collector_soiling_factor: float = 1.0,
    collector_cost_per_m2: float = 300.0,
    short_term_cost_coeff: float = 1400.0,
    short_term_cost_exponent: float = 0.78,
    seasonal_cost_coeff: float = 900.0,
    seasonal_cost_exponent: float = 0.82,
    short_term_tank_height_m: float = TANK_GEOMETRY_DEFAULTS["short_term"]["height_m"],
    seasonal_tank_height_m: float = TANK_GEOMETRY_DEFAULTS["seasonal"]["height_m"],
    short_term_static_mixing_coeff: float = TANK_GEOMETRY_DEFAULTS["short_term"]["static_mixing_coeff"],
    seasonal_static_mixing_coeff: float = TANK_GEOMETRY_DEFAULTS["seasonal"]["static_mixing_coeff"],
    short_term_charge_mixing_coeff: float = TANK_GEOMETRY_DEFAULTS["short_term"]["charge_mixing_coeff"],
    seasonal_charge_mixing_coeff: float = TANK_GEOMETRY_DEFAULTS["seasonal"]["charge_mixing_coeff"],
    short_term_discharge_mixing_coeff: float = TANK_GEOMETRY_DEFAULTS["short_term"]["discharge_mixing_coeff"],
    seasonal_discharge_mixing_coeff: float = TANK_GEOMETRY_DEFAULTS["seasonal"]["discharge_mixing_coeff"],
    electricity_price: float = 0.12,
    project_lifetime_years: float = 20.0,
    seasonal_max_years: int = 3,
    seasonal_convergence_tolerance_kwh: float = 1.0,
    return_hourly: bool = True,
) -> tuple[dict[str, float], pd.DataFrame]:
    if tank_type not in TANK_TYPE_DEFAULTS:
        raise ValueError(f"Unsupported tank type: {tank_type}")
    if tank_max_temp <= tank_min_temp:
        raise ValueError("Tank maximum temperature must be higher than tank minimum temperature.")
    if heating_delta_t < 0:
        raise ValueError("Heating system delta T must be non-negative.")
    if project_lifetime_years <= 0:
        raise ValueError("Project lifetime must be greater than zero.")

    initial_temp = tank_min_temp if initial_tank_temp is None else initial_tank_temp
    initial_temp = float(np.clip(initial_temp, tank_min_temp, tank_max_temp))
    heating_return_temp = heating_supply_temp - heating_delta_t

    capacity_kwh = tank_volume_m3 * energy_per_m3_kwh(tank_min_temp, tank_max_temp)
    num_nodes = default_stratification_nodes(tank_type)
    kwh_per_k_per_node = layer_kwh_per_k(tank_volume_m3, num_nodes)
    tank_geometry = tank_geometry_metrics(
        tank_volume_m3=tank_volume_m3,
        tank_height_m=tank_geometry_value(
            tank_type,
            short_term_value=short_term_tank_height_m,
            seasonal_value=seasonal_tank_height_m,
        ),
        static_mixing_coeff=tank_geometry_value(
            tank_type,
            short_term_value=short_term_static_mixing_coeff,
            seasonal_value=seasonal_static_mixing_coeff,
        ),
        charge_mixing_coeff=tank_geometry_value(
            tank_type,
            short_term_value=short_term_charge_mixing_coeff,
            seasonal_value=seasonal_charge_mixing_coeff,
        ),
        discharge_mixing_coeff=tank_geometry_value(
            tank_type,
            short_term_value=short_term_discharge_mixing_coeff,
            seasonal_value=seasonal_discharge_mixing_coeff,
        ),
    )
    tank_layers_c = create_tank_layers(initial_temp, num_nodes)
    seasonal_years = min(max(seasonal_max_years, 1), 3)
    if tank_type != "seasonal":
        seasonal_years = 1

    final_totals = None
    final_records = None
    converged = False
    years_used = 1

    for pass_index in range(seasonal_years):
        start_energy_kwh = tank_energy_from_layers(tank_layers_c, tank_min_temp, kwh_per_k_per_node)
        totals, records_df, end_layers = _simulate_single_pass(
            weather_load_df=weather_load_df,
            collector_area_m2=collector_area_m2,
            tank_layers_c=tank_layers_c,
            tank_type=tank_type,
            tank_min_temp=tank_min_temp,
            tank_max_temp=tank_max_temp,
            heating_supply_temp=heating_supply_temp,
            heating_delta_t=heating_delta_t,
            solar_collector_delta_t=solar_collector_delta_t,
            collector_flow_lph_per_m2=collector_flow_lph_per_m2,
            collector_delta_t_min_c=collector_delta_t_min_c,
            collector_delta_t_max_c=collector_delta_t_max_c,
            direct_solar_delivery_margin_c=direct_solar_delivery_margin_c,
            tank_thermal_loss_percent_per_day=tank_thermal_loss_percent_per_day,
            auxiliary_system=auxiliary_system,
            auxiliary_efficiency=auxiliary_efficiency,
            backup_auxiliary_efficiency=backup_auxiliary_efficiency,
            timestep_hours=timestep_hours,
            eta0=eta0,
            a1=a1,
            a2=a2,
            collector_soiling_factor=collector_soiling_factor,
            kwh_per_k_per_node=kwh_per_k_per_node,
            tank_geometry=tank_geometry,
            pass_index=pass_index,
            return_hourly=return_hourly,
        )
        end_energy_kwh = tank_energy_from_layers(end_layers, tank_min_temp, kwh_per_k_per_node)
        final_totals = totals
        final_records = records_df
        years_used = pass_index + 1
        tank_layers_c = end_layers

        if tank_type != "seasonal":
            break
        if abs(end_energy_kwh - start_energy_kwh) <= max(seasonal_convergence_tolerance_kwh, 0.005 * capacity_kwh):
            converged = True
            break

    totals = final_totals if final_totals is not None else {}
    records_df = final_records if final_records is not None else pd.DataFrame()

    collector_capital = collector_area_m2 * collector_cost_per_m2
    tank_capital = tank_capital_cost(
        tank_volume_m3=tank_volume_m3,
        tank_type=tank_type,
        short_term_cost_coeff=short_term_cost_coeff,
        short_term_cost_exponent=short_term_cost_exponent,
        seasonal_cost_coeff=seasonal_cost_coeff,
        seasonal_cost_exponent=seasonal_cost_exponent,
    )
    annualized_capital = (collector_capital + tank_capital) / project_lifetime_years
    annual_auxiliary_energy_cost = totals["auxiliary_energy_input_kwh"] * electricity_price
    annualized_total_cost = annualized_capital + annual_auxiliary_energy_cost

    productive_efficiency_pct = 0.0
    productive_efficiency_count = totals.pop("_productive_efficiency_count", 0.0)
    productive_efficiency_sum = totals.pop("_productive_efficiency_sum", 0.0)
    final_tank_temp_c = totals.pop("_final_tank_temp_c", initial_temp)
    if productive_efficiency_count > 0:
        productive_efficiency_pct = 100.0 * productive_efficiency_sum / productive_efficiency_count
    elif not records_df.empty:
        productive_efficiency_pct = 100.0 * float(records_df["collector_efficiency"].mean())

    results = {
        "tank_type": tank_type,
        "tank_type_label": tank_type_label(tank_type),
        "collector_area_m2": collector_area_m2,
        "tank_volume_m3": tank_volume_m3,
        "tank_capacity_kwh": capacity_kwh,
        "initial_tank_temp_c": initial_temp,
        "tank_final_temp_c": final_tank_temp_c,
        "tank_height_m": tank_geometry["tank_height_m"],
        "tank_diameter_m": tank_geometry["tank_diameter_m"],
        "tank_height_to_diameter_ratio": tank_geometry["tank_height_to_diameter_ratio"],
        "stratification_efficiency": tank_geometry["stratification_efficiency"],
        "static_mixing_coeff_effective": tank_geometry["static_mixing_coeff_effective"],
        "charge_mixing_coeff_effective": tank_geometry["charge_mixing_coeff_effective"],
        "discharge_mixing_coeff_effective": tank_geometry["discharge_mixing_coeff_effective"],
        "heating_return_temp_c": heating_return_temp,
        "seasonal_years_used": years_used,
        "seasonal_converged": converged if tank_type == "seasonal" else True,
        "seasonal_max_years": seasonal_years,
        "collector_cost_per_m2": collector_cost_per_m2,
        "collector_capital_cost": collector_capital,
        "tank_capital_cost": tank_capital,
        "annualized_capital_cost": annualized_capital,
        "annual_auxiliary_energy_cost": annual_auxiliary_energy_cost,
        "annualized_total_cost": annualized_total_cost,
        "collector_efficiency_pct": productive_efficiency_pct,
        **totals,
    }
    elapsed_hours = max(len(weather_load_df) * timestep_hours, 1e-9)
    average_tank_energy_above_supply_kwh = results["tank_energy_above_supply_kwh_hours"] / elapsed_hours
    average_tank_energy_between_return_supply_kwh = (
        results["tank_energy_between_return_supply_kwh_hours"] / elapsed_hours
    )
    average_tank_energy_below_return_kwh = results["tank_energy_below_return_kwh_hours"] / elapsed_hours
    average_tank_stored_energy_kwh = (
        average_tank_energy_above_supply_kwh
        + average_tank_energy_between_return_supply_kwh
        + average_tank_energy_below_return_kwh
    )
    if average_tank_stored_energy_kwh > 0:
        results["average_tank_energy_above_supply_pct"] = (
            100.0 * average_tank_energy_above_supply_kwh / average_tank_stored_energy_kwh
        )
        results["average_tank_energy_between_return_supply_pct"] = (
            100.0 * average_tank_energy_between_return_supply_kwh / average_tank_stored_energy_kwh
        )
        results["average_tank_energy_below_return_pct"] = (
            100.0 * average_tank_energy_below_return_kwh / average_tank_stored_energy_kwh
        )
    else:
        results["average_tank_energy_above_supply_pct"] = 0.0
        results["average_tank_energy_between_return_supply_pct"] = 0.0
        results["average_tank_energy_below_return_pct"] = 0.0
    results["average_tank_energy_above_supply_kwh"] = average_tank_energy_above_supply_kwh
    results["average_tank_energy_between_return_supply_kwh"] = average_tank_energy_between_return_supply_kwh
    results["average_tank_energy_below_return_kwh"] = average_tank_energy_below_return_kwh
    results["average_tank_stored_energy_kwh"] = average_tank_stored_energy_kwh
    load_total = results["heating_load_kwh"]
    results["solar_heating_ratio_pct"] = (
        100.0 * results["solar_source_energy_used_kwh"] / load_total if load_total > 0 else 0.0
    )
    results["solar_coverage_pct"] = results["solar_heating_ratio_pct"]
    results["auxiliary_fraction_pct"] = (
        100.0 * results["auxiliary_heating_delivered_kwh"] / load_total if load_total > 0 else 0.0
    )
    return results, records_df


def run_design_grid(
    weather_load_df: pd.DataFrame,
    collector_area_list: list[float],
    tank_size_list: list[float],
    tank_types: list[str] | tuple[str, ...] | None = None,
    tank_min_temp: float = 30.0,
    tank_max_temp: float = 80.0,
    initial_tank_temp: float | None = None,
    heating_supply_temp: float = 45.0,
    heating_delta_t: float = 10.0,
    solar_collector_delta_t: float = 10.0,
    collector_flow_lph_per_m2: float = 20.0,
    collector_delta_t_min_c: float = 2.0,
    collector_delta_t_max_c: float = 15.0,
    direct_solar_delivery_margin_c: float = 0.0,
    short_term_loss_percent_per_day: float = 1.0,
    seasonal_loss_percent_per_day: float = 0.15,
    auxiliary_system: str = "electric",
    auxiliary_efficiency: float = 0.98,
    backup_auxiliary_efficiency: float = 0.98,
    timestep_hours: float = 1.0,
    eta0: float = 0.72,
    a1: float = 3.2,
    a2: float = 0.0,
    collector_soiling_factor: float = 1.0,
    collector_cost_per_m2: float = 300.0,
    short_term_cost_coeff: float = 1400.0,
    short_term_cost_exponent: float = 0.78,
    seasonal_cost_coeff: float = 900.0,
    seasonal_cost_exponent: float = 0.82,
    short_term_tank_height_m: float = TANK_GEOMETRY_DEFAULTS["short_term"]["height_m"],
    seasonal_tank_height_m: float = TANK_GEOMETRY_DEFAULTS["seasonal"]["height_m"],
    short_term_static_mixing_coeff: float = TANK_GEOMETRY_DEFAULTS["short_term"]["static_mixing_coeff"],
    seasonal_static_mixing_coeff: float = TANK_GEOMETRY_DEFAULTS["seasonal"]["static_mixing_coeff"],
    short_term_charge_mixing_coeff: float = TANK_GEOMETRY_DEFAULTS["short_term"]["charge_mixing_coeff"],
    seasonal_charge_mixing_coeff: float = TANK_GEOMETRY_DEFAULTS["seasonal"]["charge_mixing_coeff"],
    short_term_discharge_mixing_coeff: float = TANK_GEOMETRY_DEFAULTS["short_term"]["discharge_mixing_coeff"],
    seasonal_discharge_mixing_coeff: float = TANK_GEOMETRY_DEFAULTS["seasonal"]["discharge_mixing_coeff"],
    electricity_price: float = 0.12,
    project_lifetime_years: float = 20.0,
    progress_callback: Callable[[dict[str, float | str]], None] | None = None,
) -> pd.DataFrame:
    rows = []
    normalized_tank_types = normalize_tank_types(tank_types)
    total = len(normalized_tank_types) * len(collector_area_list) * len(tank_size_list)
    current = 0
    for tank_type in normalized_tank_types:
        current_loss = tank_loss_percent_per_day(
            tank_type=tank_type,
            short_term_loss_percent_per_day=short_term_loss_percent_per_day,
            seasonal_loss_percent_per_day=seasonal_loss_percent_per_day,
        )
        for area in collector_area_list:
            for tank_size in tank_size_list:
                current += 1
                if progress_callback is not None:
                    progress_callback(
                        {
                            "current": current,
                            "total": total,
                            "tank_type": tank_type,
                            "collector_area_m2": area,
                            "tank_volume_m3": tank_size,
                        }
                    )
                summary, _ = simulate_design(
                    weather_load_df=weather_load_df,
                    collector_area_m2=area,
                    tank_volume_m3=tank_size,
                    tank_type=tank_type,
                    tank_min_temp=tank_min_temp,
                    tank_max_temp=tank_max_temp,
                    initial_tank_temp=initial_tank_temp,
                    heating_supply_temp=heating_supply_temp,
                    heating_delta_t=heating_delta_t,
                    solar_collector_delta_t=solar_collector_delta_t,
                    collector_flow_lph_per_m2=collector_flow_lph_per_m2,
                    collector_delta_t_min_c=collector_delta_t_min_c,
                    collector_delta_t_max_c=collector_delta_t_max_c,
                    direct_solar_delivery_margin_c=direct_solar_delivery_margin_c,
                    tank_thermal_loss_percent_per_day=current_loss,
                    auxiliary_system=auxiliary_system,
                    auxiliary_efficiency=auxiliary_efficiency,
                    backup_auxiliary_efficiency=backup_auxiliary_efficiency,
                    timestep_hours=timestep_hours,
                    eta0=eta0,
                    a1=a1,
                    a2=a2,
                    collector_soiling_factor=collector_soiling_factor,
                    collector_cost_per_m2=collector_cost_per_m2,
                    short_term_cost_coeff=short_term_cost_coeff,
                    short_term_cost_exponent=short_term_cost_exponent,
                    seasonal_cost_coeff=seasonal_cost_coeff,
                    seasonal_cost_exponent=seasonal_cost_exponent,
                    short_term_tank_height_m=short_term_tank_height_m,
                    seasonal_tank_height_m=seasonal_tank_height_m,
                    short_term_static_mixing_coeff=short_term_static_mixing_coeff,
                    seasonal_static_mixing_coeff=seasonal_static_mixing_coeff,
                    short_term_charge_mixing_coeff=short_term_charge_mixing_coeff,
                    seasonal_charge_mixing_coeff=seasonal_charge_mixing_coeff,
                    short_term_discharge_mixing_coeff=short_term_discharge_mixing_coeff,
                    seasonal_discharge_mixing_coeff=seasonal_discharge_mixing_coeff,
                    electricity_price=electricity_price,
                    project_lifetime_years=project_lifetime_years,
                    return_hourly=False,
                )
                rows.append(summary)
    return pd.DataFrame(rows).sort_values(["tank_type", "collector_area_m2", "tank_volume_m3"]).reset_index(drop=True)


def run_collector_area_study(
    weather_load_df: pd.DataFrame,
    collector_area_list: list[float],
    tank_types: list[str] | tuple[str, ...] | None = None,
    tank_min_temp: float = 30.0,
    tank_max_temp: float = 80.0,
    initial_tank_temp: float | None = None,
    heating_supply_temp: float = 45.0,
    heating_delta_t: float = 10.0,
    solar_collector_delta_t: float = 10.0,
    collector_flow_lph_per_m2: float = 20.0,
    collector_delta_t_min_c: float = 2.0,
    collector_delta_t_max_c: float = 15.0,
    direct_solar_delivery_margin_c: float = 0.0,
    short_term_loss_percent_per_day: float = 1.0,
    seasonal_loss_percent_per_day: float = 0.15,
    short_term_ratio_l_per_m2: float = 200.0,
    seasonal_ratio_l_per_m2: float = 2000.0,
    auxiliary_system: str = "electric",
    auxiliary_efficiency: float = 0.98,
    backup_auxiliary_efficiency: float = 0.98,
    timestep_hours: float = 1.0,
    eta0: float = 0.72,
    a1: float = 3.2,
    a2: float = 0.0,
    collector_soiling_factor: float = 1.0,
    collector_cost_per_m2: float = 300.0,
    short_term_cost_coeff: float = 1400.0,
    short_term_cost_exponent: float = 0.78,
    seasonal_cost_coeff: float = 900.0,
    seasonal_cost_exponent: float = 0.82,
    short_term_tank_height_m: float = TANK_GEOMETRY_DEFAULTS["short_term"]["height_m"],
    seasonal_tank_height_m: float = TANK_GEOMETRY_DEFAULTS["seasonal"]["height_m"],
    short_term_static_mixing_coeff: float = TANK_GEOMETRY_DEFAULTS["short_term"]["static_mixing_coeff"],
    seasonal_static_mixing_coeff: float = TANK_GEOMETRY_DEFAULTS["seasonal"]["static_mixing_coeff"],
    short_term_charge_mixing_coeff: float = TANK_GEOMETRY_DEFAULTS["short_term"]["charge_mixing_coeff"],
    seasonal_charge_mixing_coeff: float = TANK_GEOMETRY_DEFAULTS["seasonal"]["charge_mixing_coeff"],
    short_term_discharge_mixing_coeff: float = TANK_GEOMETRY_DEFAULTS["short_term"]["discharge_mixing_coeff"],
    seasonal_discharge_mixing_coeff: float = TANK_GEOMETRY_DEFAULTS["seasonal"]["discharge_mixing_coeff"],
    electricity_price: float = 0.12,
    project_lifetime_years: float = 20.0,
    progress_callback: Callable[[dict[str, float | str]], None] | None = None,
) -> pd.DataFrame:
    rows = []
    normalized_tank_types = normalize_tank_types(tank_types)
    total = len(normalized_tank_types) * len(collector_area_list)
    current = 0
    for tank_type in normalized_tank_types:
        current_loss = tank_loss_percent_per_day(
            tank_type=tank_type,
            short_term_loss_percent_per_day=short_term_loss_percent_per_day,
            seasonal_loss_percent_per_day=seasonal_loss_percent_per_day,
        )
        ratio_m3_per_m2 = tank_type_ratio_m3_per_m2(
            tank_type=tank_type,
            short_term_ratio_l_per_m2=short_term_ratio_l_per_m2,
            seasonal_ratio_l_per_m2=seasonal_ratio_l_per_m2,
        )
        for area in collector_area_list:
            current += 1
            if progress_callback is not None:
                progress_callback(
                    {
                        "current": current,
                        "total": total,
                        "tank_type": tank_type,
                        "collector_area_m2": area,
                        "tank_volume_m3": area * ratio_m3_per_m2,
                    }
                )
            summary, _ = simulate_design(
                weather_load_df=weather_load_df,
                collector_area_m2=area,
                tank_volume_m3=area * ratio_m3_per_m2,
                tank_type=tank_type,
                tank_min_temp=tank_min_temp,
                tank_max_temp=tank_max_temp,
                initial_tank_temp=initial_tank_temp,
                heating_supply_temp=heating_supply_temp,
                heating_delta_t=heating_delta_t,
                solar_collector_delta_t=solar_collector_delta_t,
                collector_flow_lph_per_m2=collector_flow_lph_per_m2,
                collector_delta_t_min_c=collector_delta_t_min_c,
                collector_delta_t_max_c=collector_delta_t_max_c,
                direct_solar_delivery_margin_c=direct_solar_delivery_margin_c,
                tank_thermal_loss_percent_per_day=current_loss,
                auxiliary_system=auxiliary_system,
                auxiliary_efficiency=auxiliary_efficiency,
                backup_auxiliary_efficiency=backup_auxiliary_efficiency,
                timestep_hours=timestep_hours,
                eta0=eta0,
                a1=a1,
                a2=a2,
                collector_soiling_factor=collector_soiling_factor,
                collector_cost_per_m2=collector_cost_per_m2,
                short_term_cost_coeff=short_term_cost_coeff,
                short_term_cost_exponent=short_term_cost_exponent,
                seasonal_cost_coeff=seasonal_cost_coeff,
                seasonal_cost_exponent=seasonal_cost_exponent,
                short_term_tank_height_m=short_term_tank_height_m,
                seasonal_tank_height_m=seasonal_tank_height_m,
                short_term_static_mixing_coeff=short_term_static_mixing_coeff,
                seasonal_static_mixing_coeff=seasonal_static_mixing_coeff,
                short_term_charge_mixing_coeff=short_term_charge_mixing_coeff,
                seasonal_charge_mixing_coeff=seasonal_charge_mixing_coeff,
                short_term_discharge_mixing_coeff=short_term_discharge_mixing_coeff,
                seasonal_discharge_mixing_coeff=seasonal_discharge_mixing_coeff,
                electricity_price=electricity_price,
                project_lifetime_years=project_lifetime_years,
                return_hourly=False,
            )
            summary["study_ratio_l_per_m2"] = ratio_m3_per_m2 * 1000.0
            rows.append(summary)
    return pd.DataFrame(rows).sort_values(["tank_type", "collector_area_m2"]).reset_index(drop=True)


def run_tank_size_study(
    weather_load_df: pd.DataFrame,
    collector_area_m2: float,
    tank_size_list: list[float],
    tank_types: list[str] | tuple[str, ...] | None = None,
    tank_min_temp: float = 30.0,
    tank_max_temp: float = 80.0,
    initial_tank_temp: float | None = None,
    heating_supply_temp: float = 45.0,
    heating_delta_t: float = 10.0,
    solar_collector_delta_t: float = 10.0,
    collector_flow_lph_per_m2: float = 20.0,
    collector_delta_t_min_c: float = 2.0,
    collector_delta_t_max_c: float = 15.0,
    direct_solar_delivery_margin_c: float = 0.0,
    short_term_loss_percent_per_day: float = 1.0,
    seasonal_loss_percent_per_day: float = 0.15,
    auxiliary_system: str = "electric",
    auxiliary_efficiency: float = 0.98,
    backup_auxiliary_efficiency: float = 0.98,
    timestep_hours: float = 1.0,
    eta0: float = 0.72,
    a1: float = 3.2,
    a2: float = 0.0,
    collector_soiling_factor: float = 1.0,
    collector_cost_per_m2: float = 300.0,
    short_term_cost_coeff: float = 1400.0,
    short_term_cost_exponent: float = 0.78,
    seasonal_cost_coeff: float = 900.0,
    seasonal_cost_exponent: float = 0.82,
    short_term_tank_height_m: float = TANK_GEOMETRY_DEFAULTS["short_term"]["height_m"],
    seasonal_tank_height_m: float = TANK_GEOMETRY_DEFAULTS["seasonal"]["height_m"],
    short_term_static_mixing_coeff: float = TANK_GEOMETRY_DEFAULTS["short_term"]["static_mixing_coeff"],
    seasonal_static_mixing_coeff: float = TANK_GEOMETRY_DEFAULTS["seasonal"]["static_mixing_coeff"],
    short_term_charge_mixing_coeff: float = TANK_GEOMETRY_DEFAULTS["short_term"]["charge_mixing_coeff"],
    seasonal_charge_mixing_coeff: float = TANK_GEOMETRY_DEFAULTS["seasonal"]["charge_mixing_coeff"],
    short_term_discharge_mixing_coeff: float = TANK_GEOMETRY_DEFAULTS["short_term"]["discharge_mixing_coeff"],
    seasonal_discharge_mixing_coeff: float = TANK_GEOMETRY_DEFAULTS["seasonal"]["discharge_mixing_coeff"],
    electricity_price: float = 0.12,
    project_lifetime_years: float = 20.0,
    progress_callback: Callable[[dict[str, float | str]], None] | None = None,
) -> pd.DataFrame:
    rows = []
    normalized_tank_types = normalize_tank_types(tank_types)
    total = len(normalized_tank_types) * len(tank_size_list)
    current = 0
    for tank_type in normalized_tank_types:
        current_loss = tank_loss_percent_per_day(
            tank_type=tank_type,
            short_term_loss_percent_per_day=short_term_loss_percent_per_day,
            seasonal_loss_percent_per_day=seasonal_loss_percent_per_day,
        )
        for tank_size in tank_size_list:
            current += 1
            if progress_callback is not None:
                progress_callback(
                    {
                        "current": current,
                        "total": total,
                        "tank_type": tank_type,
                        "collector_area_m2": collector_area_m2,
                        "tank_volume_m3": tank_size,
                    }
                )
            summary, _ = simulate_design(
                weather_load_df=weather_load_df,
                collector_area_m2=collector_area_m2,
                tank_volume_m3=tank_size,
                tank_type=tank_type,
                tank_min_temp=tank_min_temp,
                tank_max_temp=tank_max_temp,
                initial_tank_temp=initial_tank_temp,
                heating_supply_temp=heating_supply_temp,
                heating_delta_t=heating_delta_t,
                solar_collector_delta_t=solar_collector_delta_t,
                collector_flow_lph_per_m2=collector_flow_lph_per_m2,
                collector_delta_t_min_c=collector_delta_t_min_c,
                collector_delta_t_max_c=collector_delta_t_max_c,
                direct_solar_delivery_margin_c=direct_solar_delivery_margin_c,
                tank_thermal_loss_percent_per_day=current_loss,
                auxiliary_system=auxiliary_system,
                auxiliary_efficiency=auxiliary_efficiency,
                backup_auxiliary_efficiency=backup_auxiliary_efficiency,
                timestep_hours=timestep_hours,
                eta0=eta0,
                a1=a1,
                a2=a2,
                collector_soiling_factor=collector_soiling_factor,
                collector_cost_per_m2=collector_cost_per_m2,
                short_term_cost_coeff=short_term_cost_coeff,
                short_term_cost_exponent=short_term_cost_exponent,
                seasonal_cost_coeff=seasonal_cost_coeff,
                seasonal_cost_exponent=seasonal_cost_exponent,
                short_term_tank_height_m=short_term_tank_height_m,
                seasonal_tank_height_m=seasonal_tank_height_m,
                short_term_static_mixing_coeff=short_term_static_mixing_coeff,
                seasonal_static_mixing_coeff=seasonal_static_mixing_coeff,
                short_term_charge_mixing_coeff=short_term_charge_mixing_coeff,
                seasonal_charge_mixing_coeff=seasonal_charge_mixing_coeff,
                short_term_discharge_mixing_coeff=short_term_discharge_mixing_coeff,
                seasonal_discharge_mixing_coeff=seasonal_discharge_mixing_coeff,
                electricity_price=electricity_price,
                project_lifetime_years=project_lifetime_years,
                return_hourly=False,
            )
            rows.append(summary)
    return pd.DataFrame(rows).sort_values(["tank_type", "tank_volume_m3"]).reset_index(drop=True)


def objective_sort_config(objective: str = "cost") -> tuple[list[str], list[bool], str]:
    if objective == "cost":
        return ["annualized_total_cost", "solar_coverage_pct", "collector_area_m2", "tank_volume_m3"], [True, False, True, True], "annualized_total_cost"
    return ["auxiliary_heating_delivered_kwh", "solar_coverage_pct", "collector_area_m2", "tank_volume_m3"], [True, False, True, True], "auxiliary_heating_delivered_kwh"


def feasible_designs(summary_df: pd.DataFrame, min_solar_coverage_pct: float = 0.0) -> pd.DataFrame:
    if summary_df.empty:
        return summary_df
    if min_solar_coverage_pct <= 0:
        return summary_df.copy()
    feasible = summary_df[summary_df["solar_coverage_pct"] >= float(min_solar_coverage_pct)].copy()
    return feasible if not feasible.empty else summary_df.copy()


def near_optimal_designs(
    summary_df: pd.DataFrame,
    objective: str = "cost",
    cost_tolerance_pct: float = 0.0,
    min_solar_coverage_pct: float = 0.0,
) -> pd.DataFrame:
    candidates = feasible_designs(summary_df, min_solar_coverage_pct=min_solar_coverage_pct)
    if candidates.empty:
        return candidates
    sort_columns, ascending, primary_metric = objective_sort_config(objective)
    ordered = candidates.sort_values(by=sort_columns, ascending=ascending).reset_index(drop=True)
    tolerance_pct = max(float(cost_tolerance_pct), 0.0)
    if tolerance_pct <= 0:
        return ordered
    best_metric = float(ordered[primary_metric].iloc[0])
    metric_limit = best_metric * (1.0 + tolerance_pct / 100.0)
    near_optimal = ordered[ordered[primary_metric] <= metric_limit].copy()
    return near_optimal if not near_optimal.empty else ordered


def refined_candidate_values(values: list[float], center_value: float, refine_span_pct: float = 35.0) -> list[float]:
    sorted_values = sorted(set(float(value) for value in values))
    if len(sorted_values) <= 1:
        return sorted_values
    lower_bound = sorted_values[0]
    upper_bound = sorted_values[-1]
    span = upper_bound - lower_bound
    if span <= 0:
        return sorted_values

    center = float(np.clip(center_value, lower_bound, upper_bound))
    half_span = max(span * max(refine_span_pct, 1.0) / 200.0, span / max(2 * (len(sorted_values) - 1), 1))
    refined_low = max(lower_bound, center - half_span)
    refined_high = min(upper_bound, center + half_span)
    if refined_high <= refined_low:
        return sorted_values

    refined = np.linspace(refined_low, refined_high, len(sorted_values))
    refined_values = sorted({round(float(value), 6) for value in refined})
    if center not in refined_values:
        refined_values.append(round(center, 6))
    if lower_bound <= center <= upper_bound:
        refined_values.append(round(refined_low, 6))
        refined_values.append(round(refined_high, 6))
    return sorted(set(refined_values))


def merge_study_results(df: pd.DataFrame | None, refined_df: pd.DataFrame | None, sort_columns: list[str]) -> pd.DataFrame:
    frames = [frame for frame in [df, refined_df] if frame is not None and not frame.empty]
    if not frames:
        return pd.DataFrame()
    merged = pd.concat(frames, ignore_index=True)
    key_columns = [column for column in ["tank_type", "collector_area_m2", "tank_volume_m3"] if column in merged.columns]
    if key_columns:
        merged = merged.drop_duplicates(subset=key_columns, keep="last")
    return merged.sort_values(sort_columns).reset_index(drop=True)


def best_design(
    summary_df: pd.DataFrame,
    objective: str = "cost",
    selection_mode: str = "absolute_minimum",
    cost_tolerance_pct: float = 0.0,
    min_solar_coverage_pct: float = 0.0,
) -> pd.Series:
    candidates = near_optimal_designs(
        summary_df,
        objective=objective,
        cost_tolerance_pct=cost_tolerance_pct,
        min_solar_coverage_pct=min_solar_coverage_pct,
    )
    if candidates.empty:
        raise ValueError("No design cases are available to select from.")

    sort_columns, ascending, _ = objective_sort_config(objective)
    ordered = candidates.sort_values(by=sort_columns, ascending=ascending).reset_index(drop=True)
    if selection_mode != "compact_near_optimal" or len(ordered.index) == 1:
        return ordered.iloc[0]

    compact = ordered.sort_values(
        by=["tank_volume_m3", "collector_area_m2", sort_columns[0], "solar_coverage_pct"],
        ascending=[True, True, True, False],
    ).reset_index(drop=True)
    return compact.iloc[0]


def estimate_reference_collector_efficiency(
    weather_load_df: pd.DataFrame,
    heating_supply_temp: float,
    solar_collector_delta_t: float,
    collector_flow_lph_per_m2: float,
    collector_delta_t_min_c: float,
    collector_delta_t_max_c: float,
    eta0: float,
    a1: float,
    a2: float,
    collector_soiling_factor: float = 1.0,
    timestep_hours: float = 1.0,
) -> float:
    positive_radiation = weather_load_df[weather_load_df["solar_rad"] > 0].copy()
    if positive_radiation.empty:
        return max(min(eta0, 0.85), 0.1)

    efficiencies = []
    for row in positive_radiation.itertuples(index=False):
        _, _, efficiency = dynamic_collector_delta_t(
            area_m2=1.0,
            solar_rad_kwh_m2=row.solar_rad,
            outside_temp_c=row.outside_temp,
            inlet_temp_c=max(heating_supply_temp, 0.0),
            timestep_hours=timestep_hours,
            eta0=eta0,
            a1=a1,
            a2=a2,
            collector_soiling_factor=collector_soiling_factor,
            nominal_delta_t_c=solar_collector_delta_t,
            collector_flow_lph_per_m2=collector_flow_lph_per_m2,
            collector_delta_t_min_c=collector_delta_t_min_c,
            collector_delta_t_max_c=collector_delta_t_max_c,
            max_outlet_temp_c=heating_supply_temp + collector_delta_t_max_c,
        )
        efficiencies.append(efficiency)
    return float(np.clip(np.mean(efficiencies), 0.1, 0.85))


def analyze_input_data(
    weather_load_df: pd.DataFrame,
    heating_supply_temp: float = 45.0,
    solar_collector_delta_t: float = 10.0,
    collector_flow_lph_per_m2: float = 20.0,
    collector_delta_t_min_c: float = 2.0,
    collector_delta_t_max_c: float = 15.0,
    eta0: float = 0.72,
    a1: float = 3.2,
    a2: float = 0.0,
    collector_soiling_factor: float = 1.0,
    timestep_hours: float = 1.0,
) -> tuple[dict[str, float], pd.DataFrame, pd.DataFrame]:
    df = weather_load_df.copy()
    df["month"] = df["timestamp"].dt.to_period("M").astype(str)
    heating_hours = df["heating_load"] > 0

    annual_heating_load_kwh = float(df["heating_load"].sum())
    average_heating_load_kwh = float(df["heating_load"].mean())
    heating_season_load_kwh = float(df.loc[heating_hours, "heating_load"].sum())
    annual_tilted_solar_kwh_m2 = float(df["solar_rad"].sum())
    annual_horizontal_solar_kwh_m2 = float(df.get("horizontal_solar_rad", pd.Series(dtype=float)).fillna(0.0).sum())
    heating_season_tilted_solar_kwh_m2 = float(df.loc[heating_hours, "solar_rad"].sum())

    reference_efficiency = estimate_reference_collector_efficiency(
        weather_load_df=df,
        heating_supply_temp=heating_supply_temp,
        solar_collector_delta_t=solar_collector_delta_t,
        collector_flow_lph_per_m2=collector_flow_lph_per_m2,
        collector_delta_t_min_c=collector_delta_t_min_c,
        collector_delta_t_max_c=collector_delta_t_max_c,
        eta0=eta0,
        a1=a1,
        a2=a2,
        collector_soiling_factor=collector_soiling_factor,
        timestep_hours=timestep_hours,
    )

    monthly_df = (
        df.groupby("month", as_index=False)
        .agg(
            heating_load_kwh=("heating_load", "sum"),
            average_heating_load_kwh=("heating_load", "mean"),
            solar_rad_kwh_m2=("solar_rad", "sum"),
            horizontal_solar_rad_kwh_m2=("horizontal_solar_rad", "sum") if "horizontal_solar_rad" in df else ("solar_rad", "size"),
            mean_outside_temp_c=("outside_temp", "mean"),
        )
    )
    if "horizontal_solar_rad" not in df:
        monthly_df["horizontal_solar_rad_kwh_m2"] = np.nan

    short_area_min = heating_season_load_kwh / max(annual_scale(heating_season_tilted_solar_kwh_m2, reference_efficiency, 0.90), 1e-9)
    short_area_max = heating_season_load_kwh / max(annual_scale(heating_season_tilted_solar_kwh_m2, reference_efficiency, 0.70), 1e-9)
    long_area_min = annual_heating_load_kwh / max(annual_scale(annual_tilted_solar_kwh_m2, reference_efficiency, 1.00), 1e-9)
    long_area_max = annual_heating_load_kwh / max(annual_scale(annual_tilted_solar_kwh_m2, reference_efficiency, 0.95), 1e-9)

    summary = {
        "annual_heating_load_kwh": annual_heating_load_kwh,
        "average_heating_load_kwh": average_heating_load_kwh,
        "heating_season_load_kwh": heating_season_load_kwh,
        "annual_tilted_solar_kwh_m2": annual_tilted_solar_kwh_m2,
        "annual_horizontal_solar_kwh_m2": annual_horizontal_solar_kwh_m2,
        "reference_collector_efficiency": reference_efficiency,
        "short_term_collector_min_m2": round(max(short_area_min, 10.0), 1),
        "short_term_collector_max_m2": round(max(short_area_max, max(short_area_min, 10.0)), 1),
        "short_term_tank_min_m3": round(max(short_area_min * 0.01, 1.0), 1),
        "short_term_tank_max_m3": round(max(short_area_max * 0.30, 5.0), 1),
        "seasonal_collector_min_m2": round(max(long_area_min, 10.0), 1),
        "seasonal_collector_max_m2": round(max(long_area_max, max(long_area_min, 10.0)), 1),
        "seasonal_tank_min_m3": round(max(long_area_min * 0.30, 10.0), 1),
        "seasonal_tank_max_m3": round(max(long_area_max * 3.00, 30.0), 1),
    }

    duration_df = pd.DataFrame(
        {
            "heating_load_duration_kwh": np.sort(df["heating_load"].to_numpy())[::-1],
            "tilted_solar_duration_kwh_m2": np.sort(df["solar_rad"].to_numpy())[::-1],
            "horizontal_solar_duration_kwh_m2": np.sort(df.get("horizontal_solar_rad", pd.Series(np.zeros(len(df)))).to_numpy())[::-1],
        }
    )
    duration_df["hour_rank"] = np.arange(1, len(duration_df) + 1)

    return summary, monthly_df, duration_df


def annual_scale(solar_total_kwh_m2: float, reference_efficiency: float, target_ratio: float) -> float:
    return solar_total_kwh_m2 * max(reference_efficiency, 1e-6) * max(target_ratio, 1e-6)


def monthly_performance_from_hourly(hourly_df: pd.DataFrame) -> pd.DataFrame:
    working = hourly_df.copy()
    working["month"] = working["timestamp"].dt.to_period("M").astype(str)
    agg_map: dict[str, tuple[str, str | Callable]] = {
        "heating_load_kwh": ("heating_load_kwh", "sum"),
        "solar_energy_collected_kwh": ("solar_energy_collected_kwh", "sum"),
        "solar_source_energy_used_kwh": ("solar_source_energy_used_kwh", "sum"),
        "auxiliary_heating_delivered_kwh": ("auxiliary_heating_delivered_kwh", "sum"),
        "auxiliary_energy_input_kwh": ("auxiliary_energy_input_kwh", "sum"),
        "tank_charge_kwh": ("tank_charge_kwh", "sum"),
        "tank_discharge_total_kwh": ("tank_discharge_total_kwh", "sum"),
        "tank_loss_kwh": ("tank_loss_kwh", "sum"),
    }
    return working.groupby("month", as_index=False).agg(**agg_map)


def daily_performance_from_hourly(hourly_df: pd.DataFrame) -> pd.DataFrame:
    working = hourly_df.copy()
    working["date"] = working["timestamp"].dt.strftime("%Y-%m-%d")
    agg_map: dict[str, tuple[str, str | Callable]] = {
        "heating_load_kwh": ("heating_load_kwh", "sum"),
        "solar_energy_collected_kwh": ("solar_energy_collected_kwh", "sum"),
        "solar_direct_collector_to_load_kwh": ("solar_direct_collector_to_load_kwh", "sum"),
        "solar_from_tank_direct_kwh": ("solar_from_tank_direct_kwh", "sum"),
        "solar_from_tank_via_hp_kwh": ("solar_from_tank_via_hp_kwh", "sum"),
        "solar_heating_total_kwh": ("solar_heating_total_kwh", "sum"),
        "auxiliary_heating_delivered_kwh": ("auxiliary_heating_delivered_kwh", "sum"),
        "auxiliary_energy_input_kwh": ("auxiliary_energy_input_kwh", "sum"),
        "auxiliary_energy_for_hp_boost_kwh": ("auxiliary_energy_for_hp_boost_kwh", "sum"),
        "tank_charge_kwh": ("tank_charge_kwh", "sum"),
        "tank_discharge_total_kwh": ("tank_discharge_total_kwh", "sum"),
        "tank_loss_kwh": ("tank_loss_kwh", "sum"),
        "solar_energy_curtailed_kwh": ("solar_energy_curtailed_kwh", "sum"),
    }
    return working.groupby("date", as_index=False).agg(**agg_map)


def append_annual_summary_row(df: pd.DataFrame, label_column: str = "month", annual_label: str = "Annual Summary") -> pd.DataFrame:
    if df is None or df.empty or label_column not in df.columns:
        return df
    if df[label_column].astype(str).str.lower().eq(annual_label.lower()).any():
        return df

    annual_row: dict[str, object] = {label_column: annual_label}
    for column in df.columns:
        if column == label_column:
            continue
        series = df[column]
        if not pd.api.types.is_numeric_dtype(series):
            annual_row[column] = ""
            continue
        normalized = str(column).lower()
        if normalized.startswith("average_") or normalized.startswith("mean_") or normalized.endswith("_pct"):
            annual_row[column] = float(series.mean()) if len(series) else 0.0
        else:
            annual_row[column] = float(series.sum())

    return pd.concat([df, pd.DataFrame([annual_row])], ignore_index=True)


def export_case_report(
    output_path: str | Path,
    optimized_df: pd.DataFrame | None,
    step1_df: pd.DataFrame | None,
    step2_df: pd.DataFrame | None,
    final_summary: dict[str, float | str],
    final_hourly_df: pd.DataFrame | None,
    project_inputs: dict[str, float | str] | None = None,
    monthly_summary_df: pd.DataFrame | None = None,
    input_monthly_df: pd.DataFrame | None = None,
) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    project_inputs_df = pd.DataFrame([project_inputs]) if project_inputs else pd.DataFrame()
    final_summary_df = pd.DataFrame([final_summary])
    optimized_export_df = optimized_df if optimized_df is not None else pd.DataFrame()
    step1_export_df = step1_df if step1_df is not None else pd.DataFrame()
    step2_export_df = step2_df if step2_df is not None else pd.DataFrame()
    final_hourly_export_df = final_hourly_df if final_hourly_df is not None else pd.DataFrame()
    final_monthly_df = (
        monthly_summary_df.copy()
        if monthly_summary_df is not None and not monthly_summary_df.empty
        else (monthly_performance_from_hourly(final_hourly_export_df) if not final_hourly_export_df.empty else pd.DataFrame())
    )
    final_monthly_df = append_annual_summary_row(final_monthly_df)
    final_daily_df = (
        append_annual_summary_row(
            daily_performance_from_hourly(final_hourly_export_df),
            label_column="date",
            annual_label="Annual Summary",
        )
        if not final_hourly_export_df.empty
        else pd.DataFrame()
    )
    input_monthly_export_df = append_annual_summary_row(input_monthly_df.copy()) if input_monthly_df is not None and not input_monthly_df.empty else input_monthly_df

    def format_sheet(writer: pd.ExcelWriter, sheet_name: str, df: pd.DataFrame) -> None:
        if df is None or df.empty:
            return
        ws = writer.book[sheet_name]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = False

        header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
        header_font = Font(bold=True, color="1F1F1F")
        header_alignment = Alignment(horizontal="center", vertical="center")
        body_alignment = Alignment(vertical="center")
        thin_border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )

        currency_columns = {"annualized_total_cost", "annualized_capital_cost", "annual_auxiliary_energy_cost", "collector_capital_cost", "tank_capital_cost"}
        percent_columns = {
            "solar_coverage_pct",
            "solar_heating_ratio_pct",
            "auxiliary_fraction_pct",
            "collector_efficiency_pct",
            "solar_heat_ratio_pct",
        }
        integer_columns = {"seasonal_years_used", "seasonal_max_years", "simulation_year_index"}
        datetime_columns = {"timestamp"}

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            cell.border = thin_border

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = body_alignment
                cell.border = thin_border

        for col_idx, column_name in enumerate(df.columns, start=1):
            column_letter = get_column_letter(col_idx)
            series = df[column_name]
            max_len = max(
                len(str(column_name)),
                max((len(str(value)) for value in series.head(500).fillna("")), default=0),
            )
            ws.column_dimensions[column_letter].width = min(max(max_len + 2, 12), 26)

            if column_name in datetime_columns:
                for cell in ws[column_letter][1:]:
                    if cell.value is not None:
                        cell.number_format = "yyyy-mm-dd hh:mm"
            elif column_name in currency_columns:
                for cell in ws[column_letter][1:]:
                    if isinstance(cell.value, (int, float)):
                        cell.number_format = '"RMB" #,##0.00'
            elif column_name in percent_columns:
                for cell in ws[column_letter][1:]:
                    if isinstance(cell.value, (int, float)):
                        cell.number_format = "0.00"
            elif column_name in integer_columns:
                for cell in ws[column_letter][1:]:
                    if isinstance(cell.value, (int, float)):
                        cell.number_format = "0"
            elif pd.api.types.is_numeric_dtype(series):
                for cell in ws[column_letter][1:]:
                    if isinstance(cell.value, (int, float)):
                        cell.number_format = "#,##0.00"

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        if not project_inputs_df.empty:
            project_inputs_df.to_excel(writer, sheet_name="project_inputs", index=False)
            format_sheet(writer, "project_inputs", project_inputs_df)
        final_summary_df.to_excel(writer, sheet_name="final_summary", index=False)
        format_sheet(writer, "final_summary", final_summary_df)
        if not step1_export_df.empty:
            step1_export_df.to_excel(writer, sheet_name="step1_results", index=False)
            format_sheet(writer, "step1_results", step1_export_df)
        if not step2_export_df.empty:
            step2_export_df.to_excel(writer, sheet_name="step2_results", index=False)
            format_sheet(writer, "step2_results", step2_export_df)
        if not optimized_export_df.empty:
            optimized_export_df.to_excel(writer, sheet_name="optimization_grid", index=False)
            format_sheet(writer, "optimization_grid", optimized_export_df)
        final_monthly_df.to_excel(writer, sheet_name="monthly_summary", index=False)
        format_sheet(writer, "monthly_summary", final_monthly_df)
        if not final_daily_df.empty:
            final_daily_df.to_excel(writer, sheet_name="daily_summary", index=False)
            format_sheet(writer, "daily_summary", final_daily_df)
        if not final_hourly_export_df.empty:
            final_hourly_export_df.to_excel(writer, sheet_name="hourly_results", index=False)
            format_sheet(writer, "hourly_results", final_hourly_export_df)
        if input_monthly_export_df is not None and not input_monthly_export_df.empty:
            input_monthly_export_df.to_excel(writer, sheet_name="input_monthly_data", index=False)
            format_sheet(writer, "input_monthly_data", input_monthly_export_df)

    return output_path
